#!/usr/bin/env python3
"""Geometry QA linter for generated decks (no rendering available in this sandbox).

Flags: shapes out of slide bounds, and likely vertical text overflow in text boxes.
Usage: python3 qa.py ../00-overview.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400.0


def lint(path):
    prs = Presentation(path)
    W = prs.slide_width / EMU_IN
    H = prs.slide_height / EMU_IN
    issues = []
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            try:
                l = sh.left / EMU_IN; t = sh.top / EMU_IN
                w = sh.width / EMU_IN; h = sh.height / EMU_IN
            except Exception:
                continue
            if l < -0.02 or t < -0.02 or l + w > W + 0.05 or t + h > H + 0.05:
                issues.append(f"slide {si}: OUT-OF-BOUNDS '{sh.name}' "
                              f"[{l:.2f},{t:.2f},{w:.2f},{h:.2f}] (page {W:.2f}x{H:.2f})")
            # crude vertical overflow check for text boxes
            if sh.has_text_frame and w > 0.1 and h > 0.1:
                for p in sh.text_frame.paragraphs:
                    txt = "".join(r.text for r in p.runs)
                    if not txt:
                        continue
                    sz = max((r.font.size.pt for r in p.runs if r.font.size), default=14)
                    # estimate wrapped lines
                    char_w_in = sz * 0.52 / 72.0
                    est_w = len(txt) * char_w_in
                    import math
                    lines_para = max(1, math.ceil(est_w / max(w, 0.2)))
                    # accumulate per-paragraph; approximate
                    p._est_lines = lines_para
                total_lines = sum(getattr(p, "_est_lines", 0) for p in sh.text_frame.paragraphs)
                line_h = 0.0
                for p in sh.text_frame.paragraphs:
                    sz = max((r.font.size.pt for r in p.runs if r.font.size), default=14)
                    line_h = max(line_h, sz * 1.25 / 72.0)
                est_h = total_lines * line_h
                if est_h > h * 1.25 and total_lines > 1:
                    issues.append(f"slide {si}: TEXT-OVERFLOW? '{sh.name}' "
                                  f"est {total_lines} lines ~{est_h:.2f}in > box {h:.2f}in")
    if issues:
        print(f"\n{path}: {len(issues)} potential issue(s)")
        for i in issues:
            print("  -", i)
    else:
        print(f"{path}: geometry OK ({len(prs.slides._sldIdLst)} slides, page {W:.2f}x{H:.2f})")
    return issues


if __name__ == "__main__":
    total = 0
    for p in sys.argv[1:]:
        total += len(lint(p))
    sys.exit(1 if total else 0)
