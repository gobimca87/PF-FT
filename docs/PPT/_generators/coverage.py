#!/usr/bin/env python3
"""Coverage check: every one of the 145 ADRs must appear in the deck suite."""
import glob
import json
import os
from pptx import Presentation

HERE = os.path.dirname(__file__)
IDX = json.load(open(os.path.join(HERE, "adr_index.json")))
DECKS = sorted(glob.glob(os.path.join(HERE, "..", "*.pptx")))


def deck_text(path):
    prs = Presentation(path)
    chunks = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                chunks.append(sh.text_frame.text)
            if sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        chunks.append(cell.text)
        if slide.has_notes_slide:
            chunks.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(chunks)


def main():
    corpus = {os.path.basename(p): deck_text(p) for p in DECKS}
    allblob = "\n".join(corpus.values())
    all_adrs = [a for v in IDX["domains"].values() for a in v]
    missing = []
    for a in all_adrs:
        short = a["id"].replace("ADR-", "")
        if short not in allblob and a["id"] not in allblob:
            missing.append(short)
    print(f"Decks: {len(corpus)}  |  ADRs expected: {len(all_adrs)}  |  found: {len(all_adrs)-len(missing)}")
    for name in corpus:
        print(f"  - {name}")
    if missing:
        print(f"\nMISSING {len(missing)} ADRs:")
        print("  " + ", ".join(missing))
    else:
        print("\nALL 145 ADRs covered across the suite ✅")
    return missing


if __name__ == "__main__":
    import sys
    sys.exit(1 if main() else 0)
