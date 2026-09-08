#!/usr/bin/env python3
"""Orion dark-theme deck builder for the PFF AI ARB presentation suite.

Every deck is cloned from the uploaded Orion template so the brand theme
("Custom 1"), slide master, footer/logo and all 14 layouts are inherited verbatim.
The 20 sample slides are stripped; new slides are added on the DARK ("Black")
layout family and populated with our own text / tables / diagram shapes.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "_TEMPLATE", "Orion_Template_V1.0.pptx")
LOGOS = os.path.join(HERE, "..", "assets", "logos")

# ---- Orion "Custom 1" brand palette (dark theme) ----
BLACK   = RGBColor(0x00, 0x00, 0x00)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
CYAN    = RGBColor(0x2A, 0xCC, 0xFF)   # accent1
BLUE    = RGBColor(0x02, 0x83, 0xFF)   # dk2
PURPLE  = RGBColor(0x7F, 0x19, 0xBE)   # accent2
PINK    = RGBColor(0xFE, 0x25, 0x79)   # accent3
VIOLET  = RGBColor(0x53, 0x00, 0xDB)   # accent4
YELLOW  = RGBColor(0xFE, 0xD0, 0x38)   # accent5
GREEN   = RGBColor(0x48, 0xE8, 0x4A)   # accent6
GREY    = RGBColor(0xD5, 0xD5, 0xD5)   # lt2
MUTED   = RGBColor(0x9A, 0x9A, 0x9A)
PANEL   = RGBColor(0x14, 0x16, 0x1C)   # near-black panel fill
PANEL2  = RGBColor(0x1E, 0x22, 0x2B)

BODY_FONT = "Calibri"
HEAD_FONT = "Calibri"

# Layout names in the Orion template (Black = dark variants)
L_TITLE    = "Title - Tidal - Black"
L_AGENDA   = "Agenda - Black"
L_SECTION  = "Section Divider OI - Black"
L_CONTENT  = "Text Wide - Black (Title Only)"
L_SUBTITLE = "Text Wide With Subtitle - Black"
L_FINAL    = "Final Slide - Black"


class Deck:
    def __init__(self, template_path=TEMPLATE):
        self.prs = Presentation(template_path)
        self._strip_sample_slides()
        self._layouts = {l.name: l for l in self.prs.slide_layouts}

    # -- package bookkeeping --
    def _strip_sample_slides(self):
        sldIdLst = self.prs.slides._sldIdLst
        part = self.prs.part
        for sldId in list(sldIdLst):
            rId = sldId.get(qn("r:id"))
            sldIdLst.remove(sldId)
            try:
                part.drop_rel(rId)
            except Exception:
                pass

    def layout(self, name):
        if name not in self._layouts:
            raise KeyError(f"layout {name!r} not in template: {list(self._layouts)}")
        return self._layouts[name]

    def save(self, path):
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self.prs.save(path)
        return path

    # -- black background guarantee (in case a layout inherits otherwise) --
    def _force_black(self, slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = BLACK

    # ---------- slide constructors ----------
    def title_slide(self, title, subtitle=None, notes=None):
        s = self.prs.slides.add_slide(self.layout(L_TITLE))
        if s.shapes.title:
            s.shapes.title.text = title
        if subtitle:
            for ph in s.placeholders:
                if ph.placeholder_format.idx != 0 and ph.has_text_frame:
                    ph.text = subtitle
                    for p in ph.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.color.rgb = GREY
                    break
        self._notes(s, notes)
        return s

    def section_slide(self, title, notes=None):
        s = self.prs.slides.add_slide(self.layout(L_SECTION))
        ph = self._first_body(s)
        if ph is not None:
            ph.text = title
        else:
            self.text(s, title, 0.9, 3.0, 11.5, 1.5, size=40, color=WHITE, bold=True)
        self._notes(s, notes)
        return s

    def agenda_slide(self, title, items, notes=None):
        s = self.prs.slides.add_slide(self.layout(L_CONTENT))
        self._set_title(s, title)
        # two columns if long
        n = len(items)
        if n > 6:
            half = (n + 1) // 2
            self.bullets(s, items[:half], 0.9, 1.9, 5.7, 4.8, size=18, gap=10)
            self.bullets(s, items[half:], 6.9, 1.9, 5.7, 4.8, size=18, gap=10,
                         start=half)
        else:
            self.bullets(s, items, 0.9, 1.9, 11.5, 4.8, size=20, gap=12)
        self._notes(s, notes)
        return s

    def content_slide(self, title, subtitle=None, notes=None):
        s = self.prs.slides.add_slide(self.layout(L_CONTENT))
        self._set_title(s, title)
        if subtitle:
            self.text(s, subtitle, 0.9, 1.35, 11.5, 0.5, size=15, color=CYAN, bold=True)
        self._notes(s, notes)
        return s

    def final_slide(self, title, subtitle=None, notes=None):
        s = self.prs.slides.add_slide(self.layout(L_FINAL))
        if s.shapes.title:
            s.shapes.title.text = title
        if subtitle:
            self.text(s, subtitle, 0.9, 4.2, 11.5, 1.0, size=18, color=GREY)
        self._notes(s, notes)
        return s

    # ---------- helpers ----------
    def _first_body(self, slide):
        for ph in slide.placeholders:
            if ph.placeholder_format.idx != 0 and ph.has_text_frame:
                return ph
        return None

    def _set_title(self, slide, title):
        if slide.shapes.title:
            slide.shapes.title.text = title
        else:
            self.text(slide, title, 0.9, 0.45, 11.5, 0.9, size=30, color=WHITE, bold=True)

    def _notes(self, slide, notes):
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    def text(self, slide, text, x, y, w, h, size=16, color=WHITE, bold=False,
             italic=False, align=PP_ALIGN.LEFT, font=BODY_FONT, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        lines = text.split("\n")
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.name = font
            r.font.color.rgb = color
        return tb

    def bullets(self, slide, items, x, y, w, h, size=16, color=WHITE, gap=8,
                font=BODY_FONT, numbered=False, start=0):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        tf.margin_top = 0
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap)
            p.alignment = PP_ALIGN.LEFT
            # accent bullet marker
            marker = f"{start+i+1}. " if numbered else "▸ "
            r0 = p.add_run(); r0.text = marker
            r0.font.size = Pt(size); r0.font.name = font
            r0.font.color.rgb = CYAN; r0.font.bold = True
            r1 = p.add_run(); r1.text = item
            r1.font.size = Pt(size); r1.font.name = font
            r1.font.color.rgb = color
        return tb

    def box(self, slide, text, x, y, w, h, fill=PANEL, line=CYAN, textcolor=WHITE,
            size=13, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.25,
            align=PP_ALIGN.CENTER, font=BODY_FONT):
        sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = line; sp.line.width = Pt(line_w)
        sp.shadow.inherit = False
        tf = sp.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(4); tf.margin_right = Pt(4)
        tf.margin_top = Pt(3); tf.margin_bottom = Pt(3)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = align
        for i, line_txt in enumerate(text.split("\n")):
            pp = p if i == 0 else tf.add_paragraph()
            pp.alignment = align
            r = pp.add_run(); r.text = line_txt
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = textcolor; r.font.name = font
        return sp

    def connector(self, slide, x1, y1, x2, y2, color=CYAN, width=1.5, arrow=True):
        cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                        Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        cn.line.color.rgb = color
        cn.line.width = Pt(width)
        if arrow:
            ln = cn.line._get_or_add_ln()
            tail = ln.makeelement(qn("a:tailEnd"),
                                  {"type": "triangle", "w": "med", "len": "med"})
            ln.append(tail)
        cn.shadow.inherit = False
        return cn

    def image(self, slide, name, x, y, w=None, h=None):
        path = name if os.path.exists(name) else os.path.join(LOGOS, f"{name}.png")
        kw = {}
        if w: kw["width"] = Inches(w)
        if h: kw["height"] = Inches(h)
        return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)

    def table(self, slide, rows, x, y, w, h, col_widths=None, header=True,
              font_size=11, header_fill=PURPLE, body_fill=PANEL, alt_fill=PANEL2,
              header_color=WHITE, body_color=GREY, first_col_color=WHITE):
        nrows, ncols = len(rows), len(rows[0])
        gt = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y),
                                    Inches(w), Inches(h))
        tbl = gt.table
        # disable banded styling so our fills show
        tbl.first_row = header
        tbl.horz_banding = False
        if col_widths:
            total = sum(col_widths)
            for j, cw in enumerate(col_widths):
                tbl.columns[j].width = Emu(int(Inches(w) * cw / total))
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = tbl.cell(i, j)
                cell.margin_left = Pt(5); cell.margin_right = Pt(5)
                cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                if header and i == 0:
                    cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = body_fill if (i % 2 == 1) else alt_fill
                tf = cell.text_frame; tf.word_wrap = True
                p = tf.paragraphs[0]
                r = p.add_run(); r.text = str(val)
                r.font.size = Pt(font_size); r.font.name = BODY_FONT
                if header and i == 0:
                    r.font.bold = True; r.font.color.rgb = header_color
                else:
                    r.font.color.rgb = first_col_color if j == 0 else body_color
                    r.font.bold = (j == 0)
        return gt

    def chip_row(self, slide, logos, y, x0=0.9, x1=12.4, size=0.62, label=True,
                 label_size=9):
        """Evenly place a row of logo tiles, each fitted inside its cell
        (preserving aspect ratio), with optional labels underneath."""
        from PIL import Image
        n = len(logos)
        span = x1 - x0
        step = span / n
        cell_w = step * 0.82
        for i, item in enumerate(logos):
            name, label_txt = (item if isinstance(item, tuple) else (item, item))
            path = name if os.path.exists(name) else os.path.join(LOGOS, f"{name}.png")
            iw, ih = Image.open(path).size
            ar = iw / ih
            h = size
            w = h * ar
            if w > cell_w:                # too wide -> constrain by width
                w = cell_w
                h = w / ar
            cx = x0 + step * i + step / 2
            self.image(slide, path, cx - w/2, y + (size - h)/2, w=w, h=h)
            if label and label_txt:
                self.text(slide, label_txt, cx - step/2, y + size + 0.08, step, 0.3,
                          size=label_size, color=GREY, align=PP_ALIGN.CENTER)
