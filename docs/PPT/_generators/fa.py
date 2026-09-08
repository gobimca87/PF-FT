#!/usr/bin/env python3
"""FA enterprise light-theme deck builder (blue + white).

Cloned from the uploaded Orion_FATemplate.pptx design language: white background,
navy left sidebar carrying the FA crest, navy headings, professional blue accent
(replacing the sample's red), gold highlights, Cambria headings + Calibri body.
Client / executive grade.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "_TEMPLATE", "Orion_FATemplate.pptx")
LOGOS_NAVY = os.path.join(HERE, "..", "assets", "logos_navy")
LOGOS = os.path.join(HERE, "..", "assets", "logos")
CREST_WHITE = os.path.join(LOGOS, "fa_crest_white.png")

# ---- FA enterprise palette ----
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
NAVY       = RGBColor(0x13, 0x29, 0x4B)   # headings / primary text
NAVY_DEEP  = RGBColor(0x01, 0x1E, 0x41)   # sidebar / dark panels
BLUE       = RGBColor(0x2E, 0x75, 0xB6)   # professional accent (replaces red)
BLUE_DK    = RGBColor(0x1F, 0x5C, 0x99)
GOLD       = RGBColor(0xFF, 0xC7, 0x2C)   # sparing highlight
CARD       = RGBColor(0xF3, 0xF6, 0xFA)   # light card fill
CARD_LINE  = RGBColor(0xA9, 0xBA, 0xCE)   # card border — clearly visible
GREY       = RGBColor(0x5C, 0x6B, 0x7A)   # secondary text
MUTEDBLUE  = RGBColor(0x8F, 0xA6, 0xC4)   # muted on navy
INK        = NAVY

HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"

CX0 = 2.05      # content left
CX1 = 12.9      # content right
CW = CX1 - CX0  # content width (~10.85)


class Deck:
    def __init__(self, template_path=TEMPLATE):
        self.prs = Presentation(template_path)
        self._strip_sample_slides()
        self.layout = self.prs.slide_layouts[0]
        self._page = 0

    def _strip_sample_slides(self):
        lst = self.prs.slides._sldIdLst
        part = self.prs.part
        for sid in list(lst):
            rId = sid.get(qn("r:id"))
            lst.remove(sid)
            try:
                part.drop_rel(rId)
            except Exception:
                pass

    def save(self, path):
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self.prs.save(path)
        return path

    def set_notes(self, s, text):
        """Set speaker notes; the FA template's notes slide lacks a body
        placeholder, so inject one if python-pptx reports none."""
        if not text:
            return
        ns = s.notes_slide
        if ns.notes_text_frame is None:
            sp = parse_xml(
                '<p:sp %s>'
                '<p:nvSpPr><p:cNvPr id="10" name="Notes Placeholder"/>'
                '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
                '<p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
                '<p:spPr/>'
                '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>'
                '</p:sp>' % nsdecls('p', 'a'))
            ns.shapes._spTree.append(sp)
        ns.notes_text_frame.text = text

    # ---------- chrome ----------
    def _new(self, page=True, footer=True):
        s = self.prs.slides.add_slide(self.layout)
        bg = s.background.fill
        bg.solid(); bg.fore_color.rgb = WHITE
        # navy sidebar
        self._rect(s, 0, 0, 1.76, 7.5, NAVY_DEEP)
        # FA crest (white) centred on sidebar
        from PIL import Image
        iw, ih = Image.open(CREST_WHITE).size
        w = 0.74; h = w * ih / iw
        self.image(s, CREST_WHITE, (1.76 - w) / 2, 0.4, w=w, h=h)
        # gold rule + wordmark
        self._rect(s, 0.42, 0.42 + h + 0.18, 0.92, 0.022, GOLD)
        self.text(s, "PFF AI", 0.2, 0.42 + h + 0.26, 1.36, 0.3, size=13, color=WHITE,
                  bold=True, align=PP_ALIGN.CENTER, font=HEAD_FONT)
        self.text(s, "Enterprise Agentic AI", 0.2, 0.42 + h + 0.58, 1.36, 0.3, size=8.5,
                  color=MUTEDBLUE, align=PP_ALIGN.CENTER)
        if page:
            self._page += 1
            self.text(s, str(self._page), 0.2, 7.06, 1.36, 0.3, size=10, color=MUTEDBLUE,
                      align=PP_ALIGN.CENTER)
        if footer:
            self.text(s, "Proprietary & Confidential", CX0, 7.12, 6.0, 0.28, size=8.5,
                      color=MUTEDBLUE)
        return s

    # ---------- primitives ----------
    def _rect(self, s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1.0):
        sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = line; sp.line.width = Pt(line_w)
        sp.shadow.inherit = False
        return sp

    def text(self, s, text, x, y, w, h, size=14, color=INK, bold=False, italic=False,
             align=PP_ALIGN.LEFT, font=BODY_FONT, anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run(); r.text = line
            r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
            r.font.name = font; r.font.color.rgb = color
        return tb

    def bullets(self, s, items, x, y, w, h, size=15, color=INK, gap=9, marker_col=BLUE,
                font=BODY_FONT):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = 0; tf.margin_top = 0
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap); p.alignment = PP_ALIGN.LEFT
            r0 = p.add_run(); r0.text = "▪  "
            r0.font.size = Pt(size); r0.font.name = font
            r0.font.color.rgb = marker_col; r0.font.bold = True
            r1 = p.add_run(); r1.text = item
            r1.font.size = Pt(size); r1.font.name = font; r1.font.color.rgb = color
        return tb

    def card(self, s, x, y, w, h, fill=CARD, line=CARD_LINE, line_w=1.25):
        return self._rect(s, x, y, w, h, fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                          line_w=line_w)

    def num_circle(self, s, cx, cy, n, dia=0.56, fill=BLUE, tsize=18, tcolor=WHITE):
        """Filled circle centred at (cx, cy) with the number centred inside it."""
        self._rect(s, cx - dia / 2, cy - dia / 2, dia, dia, fill, shape=MSO_SHAPE.OVAL)
        tb = self.text(s, str(n), cx - dia / 2, cy - dia / 2, dia, dia, size=tsize,
                       color=tcolor, bold=True, align=PP_ALIGN.CENTER, font=HEAD_FONT,
                       anchor=MSO_ANCHOR.MIDDLE)
        return tb

    def box(self, s, text, x, y, w, h, style=None, size=13, bold=True,
            align=PP_ALIGN.CENTER, font=BODY_FONT,
            fill=None, line=None, textcolor=None, line_w=None, shape=None):
        # Back-compat: legacy callers pass fill/line/textcolor. In the FA light
        # theme every box renders as a light card (navy text) for an airy,
        # consistent look, except where a navy/accent emphasis style is asked for.
        if style is None:
            style = "navy" if (textcolor == WHITE and fill == NAVY_DEEP) else "light"
        if style == "navy":
            fill, line, tc = NAVY_DEEP, None, WHITE
        elif style == "accent":
            fill, line, tc = BLUE, None, WHITE
        else:
            fill, line, tc = CARD, CARD_LINE, NAVY
        sp = self.card(s, x, y, w, h, fill=fill, line=line)
        tf = sp.text_frame; tf.word_wrap = True
        tf.margin_left = Pt(5); tf.margin_right = Pt(5)
        tf.margin_top = Pt(3); tf.margin_bottom = Pt(3)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, ln in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run(); r.text = ln
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = tc; r.font.name = font
        return sp

    def connector(self, s, x1, y1, x2, y2, color=MUTEDBLUE, width=1.5, arrow=True):
        cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
        cn.line.color.rgb = color; cn.line.width = Pt(width)
        if arrow:
            ln = cn.line._get_or_add_ln()
            ln.append(ln.makeelement(qn("a:tailEnd"),
                                     {"type": "triangle", "w": "med", "len": "med"}))
        cn.shadow.inherit = False
        return cn

    def image(self, s, name, x, y, w=None, h=None):
        path = name if os.path.exists(name) else None
        if path is None:
            for base in (LOGOS_NAVY, LOGOS):
                cand = os.path.join(base, f"{name}.png")
                if os.path.exists(cand):
                    path = cand; break
        kw = {}
        if w: kw["width"] = Inches(w)
        if h: kw["height"] = Inches(h)
        return s.shapes.add_picture(path, Inches(x), Inches(y), **kw)

    # ---------- headers / slide types ----------
    def header(self, s, kicker, title, subtitle=None, tag=None, accent_tail=None):
        if kicker:
            self.text(s, kicker.upper(), CX0, 0.34, CW - 1.4, 0.3, size=11.5, color=BLUE,
                      bold=True)
        if tag:
            self.text(s, tag.upper(), CX1 - 1.5, 0.34, 1.5, 0.3, size=11.5, color=NAVY,
                      bold=True, align=PP_ALIGN.RIGHT)
        tb = s.shapes.add_textbox(Inches(CX0), Inches(0.62), Inches(CW), Inches(0.6))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = 0; tf.margin_top = 0
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = title
        r.font.size = Pt(27); r.font.bold = True; r.font.name = HEAD_FONT
        r.font.color.rgb = NAVY
        if accent_tail:
            r2 = p.add_run(); r2.text = "  " + accent_tail
            r2.font.size = Pt(27); r2.font.bold = True; r2.font.name = HEAD_FONT
            r2.font.color.rgb = BLUE
        if subtitle:
            self.text(s, subtitle, CX0, 1.24, CW, 0.4, size=13, color=GREY)

    def content_slide(self, title, subtitle=None, kicker="PFF AI · Technical Architecture",
                      tag=None, accent_tail=None, notes=None):
        s = self._new()
        self.header(s, kicker, title, subtitle, tag, accent_tail)
        self.set_notes(s, notes)
        return s

    def title_slide(self, title, title2, subtitle, kicker=None, notes=None):
        s = self._new(page=False, footer=True)
        # faint crest watermark on the right
        from PIL import Image
        iw, ih = Image.open(os.path.join(LOGOS, "fa_crest_navy.png")).size
        h = 4.6; w = h * iw / ih
        self.image(s, os.path.join(LOGOS, "fa_crest_navy.png"), 10.4, 1.45, w=w, h=h)
        if kicker:
            self.text(s, kicker.upper(), CX0, 1.75, 8.5, 0.4, size=14, color=BLUE, bold=True)
        self.text(s, title, CX0, 2.25, 8.4, 1.0, size=42, color=NAVY, bold=True, font=HEAD_FONT)
        self.text(s, title2, CX0, 3.25, 8.4, 1.0, size=42, color=BLUE, bold=True, font=HEAD_FONT)
        self._rect(s, CX0 + 0.02, 4.35, 1.7, 0.05, GOLD)
        self.text(s, subtitle, CX0, 4.6, 8.4, 1.2, size=17, color=GREY)
        self.set_notes(s, notes)
        return s

    def section_slide(self, title, subtitle=None, notes=None):
        s = self._new()
        self._rect(s, CX0, 2.7, 0.14, 1.2, GOLD)
        self.text(s, title, CX0 + 0.35, 2.75, CW - 0.4, 1.1, size=34, color=NAVY, bold=True,
                  font=HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)
        if subtitle:
            self.text(s, subtitle, CX0 + 0.35, 3.95, CW - 0.4, 0.6, size=15, color=GREY)
        self.set_notes(s, notes)
        return s

    def final_slide(self, title, subtitle=None, notes=None):
        s = self._new(page=True)
        self.text(s, title, CX0, 3.0, CW, 1.0, size=32, color=NAVY, bold=True, font=HEAD_FONT)
        self._rect(s, CX0 + 0.02, 4.05, 1.7, 0.05, GOLD)
        if subtitle:
            self.text(s, subtitle, CX0, 4.3, CW, 1.0, size=16, color=GREY)
        self.set_notes(s, notes)
        return s

    # ---------- components ----------
    def stat_cards(self, s, cards, y=2.2, h=1.7, x0=CX0, x1=CX1, gap=0.28):
        n = len(cards); w = (x1 - x0 - gap * (n - 1)) / n; x = x0
        for i, (big, small) in enumerate(cards):
            col = NAVY if i % 2 == 0 else BLUE
            self.card(s, x, y, w, h)
            self.text(s, big, x + 0.18, y + 0.22, w - 0.36, 0.7, size=26, color=col,
                      bold=True, font=HEAD_FONT)
            self.text(s, small, x + 0.18, y + 0.95, w - 0.36, h - 1.0, size=10.5, color=GREY,
                      bold=True)
            x += w + gap

    def dark_panel(self, s, title, pairs, x, y, w, h):
        self.card(s, x, y, w, h, fill=NAVY_DEEP, line=None)
        self.text(s, title.upper(), x + 0.25, y + 0.2, w - 0.5, 0.35, size=12, color=GOLD,
                  bold=True)
        yy = y + 0.72
        for k, v in pairs:
            self.text(s, k.upper(), x + 0.25, yy, w - 0.5, 0.28, size=9.5, color=GOLD, bold=True)
            self.text(s, v, x + 0.25, yy + 0.26, w - 0.5, 0.4, size=12.5, color=WHITE, bold=True)
            yy += 0.82

    def two_col(self, s, lt, litems, rt, ritems, y=1.85, size=14, gap=9):
        self.text(s, lt, CX0, y, 5.3, 0.4, size=15, color=BLUE, bold=True)
        self.bullets(s, litems, CX0, y + 0.5, 5.3, 4.6, size=size, gap=gap)
        self.text(s, rt, 7.35, y, 5.3, 0.4, size=15, color=BLUE, bold=True)
        self.bullets(s, ritems, 7.35, y + 0.5, 5.3, 4.6, size=size, gap=gap)

    def table(self, s, rows, x, y, w, h, col_widths=None, header=True, font_size=11):
        nr, nc = len(rows), len(rows[0])
        gt = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h))
        t = gt.table; t.first_row = header; t.horz_banding = False
        if col_widths:
            tot = sum(col_widths)
            for j, cw in enumerate(col_widths):
                t.columns[j].width = Emu(int(Inches(w) * cw / tot))
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.cell(i, j)
                c.margin_left = Pt(5); c.margin_right = Pt(5)
                c.margin_top = Pt(2); c.margin_bottom = Pt(2)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                if header and i == 0:
                    c.fill.solid(); c.fill.fore_color.rgb = NAVY
                else:
                    c.fill.solid(); c.fill.fore_color.rgb = CARD if i % 2 else WHITE
                p = c.text_frame.paragraphs[0]
                r = p.add_run(); r.text = str(val)
                r.font.size = Pt(font_size); r.font.name = BODY_FONT
                if header and i == 0:
                    r.font.bold = True; r.font.color.rgb = WHITE
                else:
                    r.font.color.rgb = NAVY if j == 0 else GREY
                    r.font.bold = (j == 0)
        return gt

    def chip_row(self, s, logos, y, x0=CX0, x1=CX1, size=0.62, label=True, label_size=9):
        from PIL import Image
        n = len(logos); span = x1 - x0; step = span / n; cell_w = step * 0.8
        for i, item in enumerate(logos):
            name, lab = (item if isinstance(item, tuple) else (item, item))
            path = None
            for base in (LOGOS_NAVY, LOGOS):
                cand = os.path.join(base, f"{name}.png")
                if os.path.exists(cand):
                    path = cand; break
            iw, ih = Image.open(path).size
            ar = iw / ih; h = size; w = h * ar
            if w > cell_w:
                w = cell_w; h = w / ar
            cx = x0 + step * i + step / 2
            self.image(s, path, cx - w / 2, y + (size - h) / 2, w=w, h=h)
            if label and lab:
                self.text(s, lab, cx - step / 2, y + size + 0.08, step, 0.3, size=label_size,
                          color=GREY, align=PP_ALIGN.CENTER)


# ============================================================================
# Domain-deck helpers (shared across decks 01-09), FA light theme, sidebar-aware
# ============================================================================
import json as _json

_ADRIDX = _json.load(open(os.path.join(HERE, "adr_index.json")))


def adrs_for(*domains):
    out = []
    for dm in domains:
        out.extend(_ADRIDX["domains"].get(str(dm), []))
    return out


def _short_id(adr_id):
    return adr_id.replace("ADR-", "")


def _st(status):
    return "Accepted" if status == "Accepted" else ("Proposed ◆" if status == "Proposed" else status[:8])


def _agenda_slide(self, title, items, notes=None):
    s = self.content_slide(title, kicker="Agenda")
    n = len(items)
    if n > 6:
        half = (n + 1) // 2
        self.bullets(s, items[:half], CX0, 1.95, 5.25, 4.6, size=15, gap=12)
        self.bullets(s, items[half:], 7.4, 1.95, 5.25, 4.6, size=15, gap=12)
    else:
        self.bullets(s, items, CX0, 1.95, CW, 4.6, size=16, gap=13)
    self.set_notes(s, notes)
    return s


def _pipeline(self, s, steps, y=3.0, h=1.0, x0=CX0, x1=CX1):
    n = len(steps); gap = 0.26
    w = (x1 - x0 - gap * (n - 1)) / n
    x = x0; prev = None
    for step in steps:
        if isinstance(step, str):
            self.box(s, step, x, y, w, h, style="light", size=12)
        else:
            self.card(s, x, y, w, h)
            self.text(s, step[0], x + 0.1, y + 0.16, w - 0.2, 0.5, size=12.5, color=NAVY,
                      bold=True, align=PP_ALIGN.CENTER, font=HEAD_FONT)
            self.text(s, step[1], x + 0.1, y + h - 0.42, w - 0.2, 0.4, size=10, color=GREY,
                      align=PP_ALIGN.CENTER)
        if prev is not None:
            self.connector(s, prev, y + h / 2, x, y + h / 2, color=BLUE, width=1.6)
        prev = x + w; x += w + gap


def _kv_panel(self, s, title, pairs, x, y, w, h, col=BLUE, size=12):
    self.card(s, x, y, w, h)
    self.text(s, title.upper(), x + 0.24, y + 0.18, w - 0.48, 0.35, size=12.5, color=col, bold=True)
    yy = y + 0.72
    for k, v in pairs:
        self.text(s, k, x + 0.24, yy, w - 0.48, 0.3, size=size, color=NAVY, bold=True)
        self.text(s, v, x + 0.24, yy + 0.27, w - 0.48, 0.5, size=size - 1, color=GREY)
        yy += 0.82


def _stat_cards(self, s, cards, y=2.2, h=1.8, x0=CX0, x1=CX1, gap=0.28):
    n = len(cards); w = (x1 - x0 - gap * (n - 1)) / n; x = x0
    for i, c in enumerate(cards):
        big, small = c[0], c[1]
        col = NAVY if i % 2 == 0 else BLUE
        self.card(s, x, y, w, h)
        self._rect(s, x, y + 0.28, 0.12, h - 0.56, col)
        self.text(s, big, x + 0.32, y + 0.24, w - 0.5, 0.8, size=22, color=col, bold=True, font=HEAD_FONT)
        self.text(s, small, x + 0.32, y + 1.0, w - 0.5, h - 1.05, size=11, color=GREY)
        x += w + gap


def _adr_index_slides(self, title, adrs, notes=None, rows_per_table=13):
    per_slide = rows_per_table * 2
    pages = [adrs[i:i + per_slide] for i in range(0, len(adrs), per_slide)]
    for pi, page in enumerate(pages):
        suffix = f"  ({pi + 1}/{len(pages)})" if len(pages) > 1 else ""
        s = self.content_slide(title + suffix, kicker="Decision record",
                               subtitle=f"{len(adrs)} decisions in this domain  ·  ◆ = Proposed, awaiting ARB")
        left = page[:rows_per_table]; right = page[rows_per_table:]
        for col, chunk in enumerate((left, right)):
            if not chunk:
                continue
            rows = [["ADR", "Decision", "Status"]]
            for a in chunk:
                t = a["title"] or ""
                if len(t) > 46:
                    t = t[:45] + "…"
                rows.append([_short_id(a["id"]), t, _st(a["status"])])
            x = CX0 + col * 5.55
            self.table(s, rows, x, 1.95, 5.25, min(0.32 * len(rows), 4.9),
                       col_widths=[1.0, 3.4, 1.05], font_size=8.5)
        if pi == len(pages) - 1:
            self.set_notes(s, notes)
    return pages


Deck.agenda_slide = _agenda_slide
Deck.pipeline = _pipeline
Deck.kv_panel = _kv_panel
Deck.stat_cards = _stat_cards
Deck.adr_index_slides = _adr_index_slides
